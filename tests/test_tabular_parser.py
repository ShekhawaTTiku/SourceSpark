"""All rows inform summaries while only bounded preview rows become evidence."""

from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
import pytest

from config import settings
from documents.extraction import ExtractionError, canonical_content
from documents.tabular_parser import parse_tabular


def workbook_bytes(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_csv_full_summary_includes_rows_beyond_preview() -> None:
    rows = "\n".join(f"{row},{row}" for row in range(1, 126))
    result = parse_tabular(f"Student ID,Waste (kg)\n{rows}".encode(), "csv")
    summary = result.locations["CSV, full-table summary"]
    assert "Data rows: 125" in summary
    assert "sum=7875" in summary
    assert "mean=63" in summary
    assert "Student ID: nonempty=125, missing=0; identifier" in summary
    assert "CSV, row 101" in result.locations
    assert "CSV, row 102" not in result.locations
    assert "later individual rows cannot be cited" in result.warnings[0]
    assert result.content == canonical_content(result.locations)


def test_csv_missing_values_are_not_zero_or_na_tokens() -> None:
    result = parse_tabular(b"Area,Weight (kg),Code\nA,10,NA\nB,,NULL\nC,20,N/A\n", "csv")
    assert "Weight (kg): nonempty=2, missing=1; min=10, max=20, sum=30, mean=15" in result.content
    assert "Code=NA" in result.locations["CSV, row 2"]
    assert "Weight (kg)=(missing)" in result.locations["CSV, row 3"]


def test_csv_quoted_multiline_unicode_and_record_locations() -> None:
    result = parse_tabular('Area,Note\nA,"Use labels, not guesses"\nB,"Line one\nLine two"\nC,café\n'.encode(), "csv")
    assert "Note=Use labels, not guesses" in result.locations["CSV, row 2"]
    assert "Line one\nLine two" in result.locations["CSV, row 3"]
    assert "Note=café" in result.locations["CSV, row 5"]


@pytest.mark.parametrize("encoding", ["cp1252", "utf-16"])
def test_csv_encoding_fallback_is_visible(encoding: str) -> None:
    result = parse_tabular("Area,Note\nA,café\n".encode(encoding), "csv")
    assert "café" in result.content
    assert result.warnings
    assert "UTF-16" in result.warnings[0] if encoding == "utf-16" else "Windows-1252" in result.warnings[0]


@pytest.mark.parametrize("content", [b"", b"a,b\n", b"a,b\n1,2,3\n", b"a,b\n\x00,1", b'a,b\n"broken,1'])
def test_empty_binary_and_malformed_csv_rejected(content: bytes) -> None:
    with pytest.raises(ExtractionError):
        parse_tabular(content, "csv")


def test_single_column_csv_is_supported() -> None:
    result = parse_tabular(b"Instruction\nRead the labels\nAsk the school office\n", "csv")
    assert "Data rows: 2. Columns: 1" in result.content


def test_duplicate_and_blank_headers_are_unambiguous() -> None:
    result = parse_tabular(b"A,A,A (2),\n1,2,3,4\n", "csv")
    assert "A=1; A (2)=2; A (2) (2)=3; Column 4=4" in result.content


def test_numeric_room_labels_are_identifiers_and_not_measurements() -> None:
    result = parse_tabular(b"Classroom number,Number of bins\n101,2\n102,4\n", "csv")
    assert "Classroom number: nonempty=2, missing=0; identifier" in result.content
    assert "Number of bins: nonempty=2, missing=0; min=2, max=4, sum=6, mean=3" in result.content


def test_excel_includes_sheets_and_physical_rows() -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Classrooms"
    first.append([None, None])
    first.append(["Area", "Bins"])
    first.append(["A", 12])
    first.append(["B", 8])
    notes = workbook.create_sheet("Guide")
    notes.append(["Section", "Instruction"])
    notes.append(["Labels", "Read the sorting label before disposal."])
    workbook.create_sheet("Blank")
    result = parse_tabular(workbook_bytes(workbook), "xlsx")
    assert "sum=20" in result.locations["Sheet 'Classrooms', full-table summary"]
    assert "Area=A; Bins=12" == result.locations["Sheet 'Classrooms', row 3"]
    assert "Read the sorting label" in result.locations["Sheet 'Guide', row 2"]
    assert any("'Blank' is empty" in note for note in result.warnings)


def test_missing_formula_cache_is_missing_not_zero() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Area", "Bins"])
    sheet.append(["A", 12])
    sheet.append(["B", "=4+4"])
    result = parse_tabular(workbook_bytes(workbook), "xlsx")
    assert "Bins=(missing)" in result.locations["Sheet 'Sheet', row 3"]
    assert "Bins: nonempty=1, missing=1; min=12, max=12, sum=12, mean=12" in result.content
    assert any("1 formula cells have no stored result" in warning for warning in result.warnings)


def test_formula_only_missing_row_still_counts_as_data() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Amount"])
    sheet.append(["=4+4"])
    result = parse_tabular(workbook_bytes(workbook), "xlsx")
    assert "Data rows: 1" in result.content
    assert "Amount=(missing)" in result.content


def test_cached_formula_results_are_used_without_calculating() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Area", "Bins"])
    sheet.append(["A", "=6+6"])
    original = BytesIO(workbook_bytes(workbook))
    patched = BytesIO()
    # openpyxl deliberately does not calculate/cache formulas. This small fixture
    # inserts a stored value to represent a workbook saved by a calculation app.
    with ZipFile(original) as source, ZipFile(patched, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<f>6+6</f><v></v>", b"<f>6+6</f><v>12</v>")
            target.writestr(item, data)
    result = parse_tabular(patched.getvalue(), "xlsx")
    assert "Bins=12" in result.locations["Sheet 'Sheet', row 2"]
    assert any("cached values may be stale" in warning for warning in result.warnings)
    assert not any("no stored result" in warning for warning in result.warnings)


def test_excel_numeric_summary_uses_all_125_rows() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Area", "Waste kg"])
    for row in range(1, 126):
        sheet.append([f"Area {row}", row])
    result = parse_tabular(workbook_bytes(workbook), "xlsx")
    assert "sum=7875" in result.content
    assert "Sheet 'Sheet', row 102" not in result.locations


def test_csv_table_cell_guard(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "MAX_TABLE_CELLS", 4)
    with pytest.raises(ExtractionError, match="table-cell extraction limit"):
        parse_tabular(b"Area,Bins\nA,12\nB,8", "csv")


def test_excel_table_cell_guard(monkeypatch: pytest.MonkeyPatch) -> None:
    workbook = Workbook()
    workbook.active.append(["Area", "Bins"])
    workbook.active.append(["A", 12])
    monkeypatch.setattr(settings, "MAX_TABLE_CELLS", 2)
    with pytest.raises(ExtractionError, match="table-cell extraction limit"):
        parse_tabular(workbook_bytes(workbook), "xlsx")


@pytest.mark.parametrize("file_type,content", [("xlsx", b"not a workbook"), ("xls", b"ignored")])
def test_corrupt_or_unsupported_tabular_file(file_type: str, content: bytes) -> None:
    with pytest.raises(ExtractionError):
        parse_tabular(content, file_type)
