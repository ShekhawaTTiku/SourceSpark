"""Source validation dispatches one real parser and preserves content identity."""

from hashlib import sha256
from io import BytesIO

from docx import Document
from openpyxl import Workbook
import pymupdf
import pytest

from config.settings import MAX_FILE_BYTES
from documents.extraction import ExtractionError
from documents.file_router import parse_uploaded_file


def test_uppercase_extension_identity_and_no_raw_binary() -> None:
    raw = b"Area,Bins\nA,12\n"
    record = parse_uploaded_file("School.CSV", raw)
    assert record.document_id == sha256(raw).hexdigest()
    assert record.filename == "School.CSV"
    assert record.file_type == "csv"
    assert record.size_bytes == len(raw)
    assert record.character_count == len(record.content)
    assert "file_bytes" not in record.model_dump()


def test_router_supports_all_four_real_formats() -> None:
    word = Document()
    word.add_paragraph("Fictional Greenfield School waste report")
    word_buffer = BytesIO()
    word.save(word_buffer)
    excel = Workbook()
    excel.active.append(["Area", "Bins"])
    excel.active.append(["A", 12])
    excel_buffer = BytesIO()
    excel.save(excel_buffer)
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((72, 72), "Fictional Greenfield School waste report")
        pdf_raw = pdf.tobytes()
    for filename, data in [("school.docx", word_buffer.getvalue()), ("school.xlsx", excel_buffer.getvalue()), ("school.pdf", pdf_raw), ("school.csv", b"Area,Bins\nA,12\n")]:
        assert parse_uploaded_file(filename, data).locations


def test_word_blank_edge_cells_before_paragraph_preserve_canonical_content() -> None:
    word = Document()
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Fictional first cell"
    table.cell(1, 1).text = "Fictional last cell"
    word.add_paragraph("A fictional school observation after the table.")
    buffer = BytesIO()
    word.save(buffer)
    record = parse_uploaded_file("school.docx", buffer.getvalue())
    assert record.locations["Table 1, row 1"] == "Fictional first cell |"
    assert record.locations["Table 1, row 2"] == "| Fictional last cell"
    assert "A fictional school observation" in record.content
    assert record.content == "\n\n".join(f"[{location}]\n{text}" for location, text in record.locations.items())


@pytest.mark.parametrize("filename", ["", "../file.csv", "C:\\secret.csv", "/tmp/file.csv", "report.csv:other", "file\x00.csv", "bad?.csv", "file.exe", "source.csv "])
def test_unsafe_or_unsupported_filename(filename: str) -> None:
    with pytest.raises(ExtractionError):
        parse_uploaded_file(filename, b"Area,Bins\nA,12\n")


def test_empty_and_oversize_source_guards() -> None:
    with pytest.raises(ExtractionError, match="empty"):
        parse_uploaded_file("school.csv", b"")
    with pytest.raises(ExtractionError, match="per-file"):
        parse_uploaded_file("school.csv", b"x" * (MAX_FILE_BYTES + 1))
