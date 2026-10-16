"""Recreate the deliberately fictional learning fixtures using project libraries.

These are small parser examples, not real statistics or advice. Generated source
files can be uploaded manually to the app and used without any Gemini request.
"""

from pathlib import Path
import csv

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
import pymupdf


def generate_samples(directory: Path | None = None) -> list[Path]:
    directory = directory or Path(__file__).resolve().parent
    directory.mkdir(parents=True, exist_ok=True)
    word_path = directory / "fictional_school_guide.docx"
    document = Document()
    document.add_heading("Fictional Greenfield School Waste Guide", level=0)
    document.add_paragraph("This is invented sample material for learning and software testing. It is not real waste-management evidence.")
    document.add_paragraph("The fictional school has 20 classrooms. Twelve classrooms have labelled sorting bins; eight do not.")
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Source detail"
    table.rows[0].cells[1].text = "Fictional finding or instruction"
    facts = [
        ("Sorting labels", "Read the displayed sorting label before putting an item into a bin."),
        ("Missing labels", "The school office supplies replacement labels to classroom monitors."),
        ("Collection information", "Collection information appears on the school noticeboard."),
        ("Uncertain items", "Students ask a classroom monitor when the displayed instructions do not identify an item."),
        ("Canteen", "The canteen displays a sorting guide beside its exit."),
        ("Review", "Classroom monitors record missing or damaged labels each Friday."),
    ]
    for label, fact in facts:
        row = table.add_row().cells
        row[0].text, row[1].text = label, fact
    document.save(word_path)

    csv_path = directory / "fictional_weekly_audit.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.writer(output)
        writer.writerow(["Audit ID", "Fictional area", "Recorded waste (kg)", "Note"])
        for number in range(1, 126):
            writer.writerow([number, f"School area {number}", number, "Invented software-testing value"])

    excel_path = directory / "fictional_school_audit.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Classrooms"
    sheet.append(["Fictional section", "Classrooms", "Labelled bins"])
    sheet.append(["Lower school", 10, 7])
    sheet.append(["Upper school", 10, 5])
    notes = workbook.create_sheet("Guide")
    notes.append(["Section", "Fictional instruction"])
    for label, fact in facts:
        notes.append([label, fact])
    notes.append(["Disclaimer", "All entries are invented for learning and testing."])
    for page in workbook:
        page.freeze_panes = "A2"
        for cell in page[1]:
            cell.font = Font(bold=True)
        page.column_dimensions["A"].width = 26
        page.column_dimensions["B"].width = 90 if page.title == "Guide" else 20
        page.column_dimensions["C"].width = 20
    workbook.save(excel_path)

    pdf_path = directory / "fictional_school_report.pdf"
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((50, 55), "Fictional Greenfield School Waste Report", fontsize=17)
        lines = [
            "Invented learning and testing material. These are not real statistics.",
            "The school has 20 classrooms; 12 have labelled sorting bins and 8 do not.",
            "The lower school has 10 classrooms and 7 labelled sorting bins.",
            "The upper school has 10 classrooms and 5 labelled sorting bins.",
            "The school office supplies replacement labels to classroom monitors.",
            "Collection information appears on the school noticeboard.",
        ]
        page.insert_textbox(pymupdf.Rect(50, 90, 550, 750), "\n\n".join(lines), fontsize=11)
        page = pdf.new_page()
        page.insert_text((50, 55), "Fictional sorting instructions", fontsize=17)
        page.insert_textbox(pymupdf.Rect(50, 90, 550, 750), "\n\n".join(fact for _, fact in facts), fontsize=11)
        pdf.save(pdf_path)
    return [word_path, csv_path, excel_path, pdf_path]


if __name__ == "__main__":
    generate_samples()
